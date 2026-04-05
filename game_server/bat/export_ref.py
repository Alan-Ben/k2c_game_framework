# -*- coding: UTF-8 -*-
import os
import openpyxl as xl

def getLimitColum(sheet):
    limitColum = 0
    for i in range(1, sheet.max_column + 1):
        title = sheet.cell(row=1, column=i).value
        title2 = sheet.cell(row=2, column=i).value
        if (title is None or len(str(title).strip()) == 0) and (title2 is None or len(str(title2).strip()) == 0):
            break
        else:
            limitColum = i
    return limitColum


# 计算跳过的列index
def calSkippedColums(row, limitColum):
    ret = []
    for i in range(limitColum):
        value = row[i].value
        if value is None or len(str(value).strip()) == 0:
            ret.append(i)
            continue
        if str(value)[0] == '#':
            ret.append(i)
            continue
    return ret


def processSheet(sheet):
    print("sheet:%s" % sheet.title)
    limitColum = getLimitColum(sheet)
    rows = list(sheet.rows)
    skipedColums = calSkippedColums(rows[1], limitColum)
    lines = []
    for row in rows[1:]:
        if row[0].value is None or len(str(row[0])) == 0:
            continue
        line = []
        for i in range(limitColum):
            if i in skipedColums:
                continue
            var = row[i].value
            if var is None:
                line.append('')
            else:
                line.append(str(var).strip())
        lines.append("\t".join(line))
    outFile = open(os.path.join(outDir, "%s.txt" % sheet.title), "wb")
    outFile.write('\n'.join(lines).encode('utf-8'))
    outFile.close()



def processFile(fileFullPathName):
    print("process file:" + fileFullPathName)

    workbook = xl.load_workbook(fileFullPathName, data_only=True)
    for sheet in workbook.worksheets:
        if sheet.max_row < 2:
            continue
        processSheet(sheet)
    workbook.close()


def walkFile(path_dir):
    for root, dirs, files in os.walk(path_dir):
        for f in files:
            if not f.endswith(".xlsx"):
                continue
            fullPath = os.path.join(root, f)
            processFile(fullPath)


outDir = "./__dlserverref"
srcDir = "./"

if __name__ == '__main__':
    os.makedirs(outDir, exist_ok=True)
    print("out dir:"+os.path.abspath(outDir))
    walkFile(srcDir)
